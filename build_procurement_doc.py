# -*- coding: utf-8 -*-
"""采购需求文档 v6（多标的版）：
- 7 类设备拆分为 7 个独立标的（表格 7 个数据行），每行配正确计量单位
- 技术参数仅列硬件规格，不含功能描述 / 使用场景 / 录像回放
- 以基础下限为标准（4 核 / 16G / 显存 4GB / 1080p / 802.3af / IP66）
- 摄像头仅支持实时查看（非 AI、无录像回放）
- 每条参数独立成行，条目间空行分隔，沿用"编号 + 参数：+ ▲ + 投标备注"行文风格
- 每个标的参数独立从 1 起编号；▲ 在编号前标记实质性要求
- 商务要求独立成表外章节
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"D:/MengWork/Web/RobotsSystem/采购需求_基础设施参数（30机器人）.docx"

def set_run_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)

def set_default_font(doc):
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')
    rFonts.set(qn('w:ascii'), '微软雅黑')
    rFonts.set(qn('w:hAnsi'), '微软雅黑')

def add_para(container, text, bold=False, size=10.5, align=None, name="微软雅黑"):
    p = container.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    set_run_font(run, name=name, size=size, bold=bold)
    return p

def merge_full(row, ncols):
    c = row.cells[0]
    for j in range(1, ncols):
        c = c.merge(row.cells[j])
    return c

# =================== 各标的（类别）技术参数（仅硬件，无商务条款） ===================
# 每个元组：(标的名称, 数量, 计量单位, [参数行...])
# 参数行格式："N. 名称：值" 或 "▲N. 名称：值"（▲ 在编号前，标记实质性要求）
CATEGORIES = [
    ("调度一体工作站（主机 / 显卡 / 存储）", 1, "台", [
        "1. CPU：≥ 4 物理核心（x86-64 架构，主频 ≥ 2.0GHz）。",
        "2. 内存：≥ 16GB DDR4。",
        "3. 系统盘：≥ 256GB SSD。",
        "4. 数据盘：≥ 512GB SSD。",
        "▲5. 显卡：独立显卡，显存 ≥ 4GB（投标时提供原厂技术规格书）。",
        "6. 网络接口：≥ 1× 千兆以太网口。",
        "▲7. 视频输出：≥ 2× 视频接口（HDMI / DP），支持双屏 1920×1080 输出（投标时提供设备检测报告）。",
        "8. 机箱与电源：塔式或机架式机箱，含电源。",
    ]),
    ("显示系统（双屏）", 1, "套", [
        "1. 显示器数量：≥ 2 台。",
        "▲2. 分辨率：≥ 1920×1080（投标时提供设备检测报告）。",
        "3. 屏幕尺寸：≥ 24 英寸。",
        "4. 面板类型：IPS。",
        "5. 视频接口：≥ 1× HDMI + ≥ 1× DP。",
        "6. 安装配件：含支架或壁挂件及连接线缆。",
    ]),
    ("监控摄像头系统（实时查看）", 1, "套", [
        "1. 摄像机数量：4 台。",
        "▲2. 分辨率：≥ 1920×1080（投标时提供设备检测报告）。",
        "3. 传感器类型：≥ 1/2.8\" CMOS 图像传感器。",
        "4. 镜头焦距：2.8mm–12mm 电动变焦。",
        "5. 视频接口：RJ45 10/100Mbps 以太网口。",
        "6. 视频编码：H.264 / H.265（硬件编码）。",
        "▲7. 供电方式：支持 PoE（IEEE 802.3af）或 DC 12V（投标时提供供电规格说明）。",
        "▲8. 防护等级：≥ IP66（投标时提供防护等级检测报告）。",
        "▲9. 工作温度范围：-10℃ ~ +50℃（投标时提供工作温度检测报告）。",
        "10. 外壳材质：铝合金或工程塑料。",
        "11. 安装接口：吸顶或壁装（含支架接口）。",
        "12. 红外补光：支持夜视（IR），补光距离 ≥ 10m。",
    ]),
    ("核心路由器", 1, "台", [
        "1. WAN 口：≥ 1× 千兆以太网口。",
        "2. LAN 口：≥ 4× 千兆以太网口。",
        "▲3. NAT 吞吐：≥ 1Gbps（投标时提供设备技术规格书）。",
        "4. 电源：外置电源适配器，DC 12V。",
    ]),
    ("PoE 交换机", 1, "台", [
        "1. 端口：≥ 8 口千兆以太网（含 PoE 口 ≥ 8）。",
        "▲2. PoE 供电：符合 IEEE 802.3af（单口 ≥ 15.4W），整机 PoE 输出 ≥ 120W（投标时提供 PoE 供电规格说明）。",
        "3. 上联端口：≥ 1× 千兆 SFP / RJ45 上联。",
        "4. 交换容量：≥ 16Gbps。",
        "5. 电源：内置电源，AC 100–240V。",
    ]),
    ("无线接入点 AP", 3, "台", [
        "1. 无线标准：IEEE 802.11ax（WiFi 6）。",
        "2. 频段：2.4GHz + 5GHz 双频。",
        "3. 并发终端：≥ 50 台。",
        "▲4. 供电方式：支持 PoE（IEEE 802.3af）（投标时提供供电规格说明）。",
        "5. 覆盖半径：≥ 15m（空旷环境）。",
        "6. 安装方式：吸顶式，含安装件。",
    ]),
    ("网络机柜及综合布线", 1, "套", [
        "1. 机柜：9–12U 壁挂式网络机柜（≥ 450×400mm 深）。",
        "2. 线缆：六类非屏蔽双绞线（Cat6）及水晶头若干。",
        "3. 电源：机柜 PDU 电源插排 × 1。",
        "4. 布线：采用直连理线，不配置独立配线架。",
    ]),
]

NOTE = ("备注：以上参数为各设备基础硬件规格下限，供应商可超配但不得低于上述下限；各设备均须符合相关国家及行业标准。")

# 商务要求（独立章节，不写入技术参数表格）
BIZ_SECTION = [
    ("（一）实施要求",
     "1、供应商须严格按照项目需求提供高质量的设备供货、安装调试与系统集成服务，确保系统按期投入运行。"
     "2、设备须与现有系统对接联调。"
     "3、供应商须制定详细的供货与安装实施方案，明确到货、上架、布线、联调、培训等阶段任务、时间节点及负责人，并定期向采购人汇报进度。"
     "4、供应商须指派一名项目经理全程负责项目协调与管理，并配备专业团队（网络工程师、系统集成工程师与现场实施人员）。"
     "5、未经采购人书面同意，供应商不得将项目分包或转包给第三方。"
     "6、项目成果的知识产权归采购人所有；项目实施涉及的系统配置、网络拓扑等信息不得泄露。"
     "7、系统建设须达到约定的并发调度指标，并完成全部设备供货、安装与联调。"),
    ("（二）服务要求",
     "1、供应商须确保所有设备符合国家及行业标准，并提供原厂质保。"
     "2、供应商须保证其提供的设备不侵犯任何第三方的知识产权。"
     "3、供应商须提供专门的售后服务人员，系统运行问题及时解决；采购人提出服务请求后 24 小时内响应，重大故障 4 小时内现场处置。"
     "4、若供应商未按合同约定提供服务，采购人有权要求赔偿。"),
    ("（三）付款及交付",
     "1、付款条件：项目款项分两阶段支付——合同签订后支付 30% 作为首付款，项目验收合格后支付剩余 70%。"
     "2、交货时间和地点：主要设备须在合同签订后 2 个月内到货并完成安装调试；交货地点为采购人指定地点。"
     "3、采购方式：本项目按政府采购相关规定采用公开招标方式采购（或按采购人批准的采购方式执行）。"
     "4、验收标准：所有设备须符合本需求表及合同要求；采购人有权对不合格部分要求整改，直至验收通过。"),
    ("（四）核心产品",
     '本项目核心产品为“调度一体工作站”（对应标的序号 1）。'),
    ("（五）特殊说明",
     "本项目不接受进口产品竞标，如竞标供应商采用进口产品竞标则作无效响应处理。"),
    ("（六）其他说明",
     "竞标供应商可结合项目实际提供技术方案、安装调试实施方案。"),
]

# =================== 构建 ===================
doc = Document()
set_default_font(doc)
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)

add_para(doc, "项目采购需求", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "说明：")
add_para(doc, "1. 本需求表及商务要求中带有“▲”的技术参数或要求为实质性要求，必须满足，响应内容不得低于该技术指标或要求。")
add_para(doc, "2. 表中各标的参数为基础硬件规格下限，供应商可超配但不得低于上述下限；各设备均须符合相关国家及行业标准。")
add_para(doc, "采购预算：________元（具体金额以采购公告 / 招标文件为准）")

# ---- 一、技术参数及规格要求（多标的，每行一标的） ----
NCOLS = 5
table = doc.add_table(rows=1, cols=NCOLS)
table.style = 'Table Grid'

title_row = table.rows[0]
tc = merge_full(title_row, NCOLS)
add_para(tc, "一、采购标的技术参数及规格要求", bold=True, size=11)

hdr = table.add_row()
for i, h in enumerate(["序号", "标的名称", "技术参数及规格要求", "数量", "单位"]):
    add_para(hdr.cells[i], h, bold=True, size=10.5)
    hdr.cells[i].width = Cm([0.9, 3.0, 10.6, 1.2, 1.3][i])

for idx, (name, count, unit, items) in enumerate(CATEGORIES, start=1):
    row = table.add_row()
    row.cells[0].width = Cm(0.9)
    row.cells[1].width = Cm(3.0)
    row.cells[2].width = Cm(10.6)
    row.cells[3].width = Cm(1.2)
    row.cells[4].width = Cm(1.3)
    add_para(row.cells[0], str(idx))
    add_para(row.cells[1], name)
    spec_cell = row.cells[2]
    for it in items:
        add_para(spec_cell, it, size=10)
        add_para(spec_cell, "", size=4)
    add_para(row.cells[3], str(count))
    add_para(row.cells[4], unit)

# ---- 二、商务要求（独立于技术参数表格之外） ----
add_para(doc, "", size=4)
sp = doc.add_paragraph()
sp.paragraph_format.space_before = Pt(6)
r = sp.add_run("▲二、商务要求")
set_run_font(r, size=11, bold=True)
for label, text in BIZ_SECTION:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label)
    set_run_font(run, size=10.5, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    run2 = p2.add_run(text)
    set_run_font(run2, size=10)

# ---- 三、其他 ----
sp3 = doc.add_paragraph()
sp3.paragraph_format.space_before = Pt(6)
r3 = sp3.add_run("三、涉及项目的其他要求及说明")
set_run_font(r3, size=11, bold=True)
p3 = doc.add_paragraph()
run3 = p3.add_run("竞标供应商须结合现场实际情况，提供完整的设备清单、拓扑设计与安装调试方案；"
                  "本需求未列明但为保障系统完整运行所必需的辅材（如电源线、理线器、固定件等），由供应商在投标方案中一并考虑，不单独计价。")
set_run_font(run3, size=10)

doc.save(OUT)
print("SAVED:", OUT, "| 标的数:", len(CATEGORIES))
