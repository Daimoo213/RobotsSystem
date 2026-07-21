# -*- coding: utf-8 -*-
"""生成《技术方案总览：功能维度·技术选型·技术路线》本地 DOCX。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\MengWork\Web\RobotsSystem\技术方案总览_功能维度_技术选型_技术路线.docx"

CJK = "微软雅黑"
NAVY = RGBColor(0x1F, 0x3B, 0x73)
BLUE = RGBColor(0x2F, 0xD7, 0xFF)
DARK = RGBColor(0x22, 0x2B, 0x35)
GREY = RGBColor(0x5A, 0x7A, 0x92)

doc = Document()

# ---- 默认字体（含中文）----
def set_cjk(style_or_run, name=CJK):
    try:
        style_or_run.font.name = name
        rpr = style_or_run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:eastAsia'), name)
        rfonts.set(qn('w:ascii'), name)
        rfonts.set(qn('w:hAnsi'), name)
    except Exception:
        pass

normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
set_cjk(normal)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def set_cell_font(cell, size=9.5, color=DARK, bold=False, cjk=True):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            if cjk:
                set_cjk(r)

def add_heading(text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.color.rgb = NAVY if level <= 1 else BLUE
    run.font.bold = True
    set_cjk(run)
    if level == 1:
        run.font.size = Pt(16)
    elif level == 2:
        run.font.size = Pt(13)
    else:
        run.font.size = Pt(11.5)
    return h

def add_para(text, size=10.5, color=DARK, bold=False, italic=False, space_after=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    set_cjk(run)
    return p

def add_bullet(text, size=10, color=DARK, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r1 = p.add_run(bold_lead)
        r1.font.bold = True
        r1.font.size = Pt(size)
        r1.font.color.rgb = NAVY
        set_cjk(r1)
    r2 = p.add_run(text)
    r2.font.size = Pt(size)
    r2.font.color.rgb = color
    set_cjk(r2)
    return p

def add_table(headers, rows, widths=None, header_fill="1F3B73"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = htext
        shade(hdr[i], header_fill)
        set_cell_font(hdr[i], size=9.5, color=RGBColor(0xFF,0xFF,0xFF), bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            set_cell_font(cells[i], size=9.5)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Cm(w)
    return t

def hrule():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2FD7FF')
    pbdr.append(bottom)
    pPr.append(pbdr)

# ===================== 封面标题 =====================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("机器人集群智能调度系统")
r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = NAVY; set_cjk(r)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("技术方案总览 · 功能维度 / 技术选型 / 技术路线")
r.font.size = Pt(14); r.font.color.rgb = BLUE; set_cjk(r)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("面向开发团队 · 宏观架构视角 · 便于快速理解整体技术方案")
r.font.size = Pt(10.5); r.font.color.rgb = GREY; set_cjk(r)
doc.add_paragraph()

# ===================== 说明 =====================
add_para("本文档从三个板块系统梳理项目的技术方案：① 功能维度——系统需要承载哪些能力；"
         "② 技术选型——在宏观层面选择了哪些架构方向与技术策略，及其概念性理由；"
         "③ 技术路线——分阶段的实施路径与里程碑。全文聚焦“做什么、选什么方向、按什么顺序做”，"
         "不涉及具体代码实现。", size=10.5, color=DARK, space_after=8)
hrule()

# ===================== 板块一：功能维度 =====================
add_heading("板块一 · 功能维度", level=1)
add_para("功能维度回答“系统要做什么”。按承载位置分为后端（系统大脑）、前端（双大屏）、"
         "机器人（设备接入）三个维度。每个功能点给出概念性解释：它是什么、解决什么问题。",
         size=10.5, color=GREY, space_after=8)

add_heading("1.1 后端功能维度（系统的“大脑”）", level=2)
add_para("后端是调度系统的核心，负责把施工计划转化为可执行的设备派单与路径，并汇聚全部运行状态。",
         size=10, color=GREY, space_after=4)
add_bullet("统一纳管 20–100 台异构施工设备（AGV、挖掘机、吊车、砌筑机器人、巡检机器人等）。"
           "每台设备建立“一机一档”（型号、能力、标段权限），通过心跳与遥测实时掌握在线态、位置、电量、任务态。",
           bold_lead="设备接入与集群管理：")
add_bullet("以网格+高度层表达微缩工地沙盘，划分作业区/禁行区/堆放区/停靠区，标注基坑、材料区、吊装位等施工点位"
           "并绑定工序与适配设备；随施工进度动态更新，并实时标注临时堆放、人员等动态障碍。",
           bold_lead="工地地图与环境感知：")
add_bullet("系统核心。将 29 项施工工序抽象为带前后置约束的任务图，依据进度计划自动编排；"
           "按“成本+优先级”把任务派给最合适的设备，支持批量派单、任务模板、人工干预（暂停/改派/插单）。",
           bold_lead="集群智能调度：")
add_bullet("把调度决策转译为设备可执行的指令：对同一标段/同型设备批量下发群组指令，按工艺节奏对齐多机节拍（编队），"
           "以及设备间实时交换位姿、进入安全距离触发减速/停车（碰撞联锁）。",
           bold_lead="集群协同作业控制：")
add_bullet("建立每台设备的能耗模型，预测剩余续航，在空闲或低优先任务间隙错峰调度充电；"
           "按运行时长/里程触发保养工单，降低损耗。",
           bold_lead="能耗与维保管控：")
add_bullet("聚合所有模块的状态与事件，驱动数字孪生大屏：全景镜像沙盘态势，支持单台设备下钻查看轨迹/电量/任务/历史告警；"
           "对离线、超时、低电量、路径冲突、越界等分级推送告警并关联处置闭环。",
           bold_lead="监控与告警中心：")
add_bullet("把调度与执行全过程沉淀为时序数据，按标段/日期/工序/车型/班组多维度聚合，"
           "一键导出标准化报表（Excel/PDF），支撑复盘与策略优化。",
           bold_lead="施工数据与报表统计：")
add_bullet("安全底线。基于实时位姿与电子围栏实现人机避让（危险区降速/停机）；"
           "支持一键全局急停；按标段划分电子围栏，跨区作业需授权，避免并行施工干扰。",
           bold_lead="施工现场安全管控：")

add_heading("1.2 前端功能维度（双大屏）", level=2)
add_para("前端是系统的“面孔”，分 PM 端与 O&M 端两套独立大屏，共用同一后端。",
         size=10, color=GREY, space_after=4)
add_bullet("面向项目经理/参观讲解。三栏布局：左栏（各类机器人投入分布、设备状态、资源负载、7 大施工阶段总览）；"
           "中栏（任务调度甘特图，支持日/周/月/季度切换、依赖线、里程碑、今日线）；右栏（待办任务与新增、摄像头视角+AI 识别浮窗、"
           "任务完成趋势三线、实时事件流）。顶部含系统时间、设备在线率、扬尘、工期健康分、8 项核心 KPI。",
           bold_lead="PM 端（施工进度指挥）：")
add_bullet("面向运维/现场调度。三栏+底部设备条：左栏（设备状态分布、8 个运维快捷指令、实时运行日志）；"
           "中栏（基于 Three.js 的三维土方调度视角，设备实时标记、场景工具栏全局指令）；右栏（集群协同拓扑、设备类型分布、实时告警、系统概览）；"
           "底部 18 张设备状态卡片（含连/定/电/任/安 5 项体检指标）。顶部含在线/故障/健康度统计与调度/能耗/安全三视图切换。",
           bold_lead="O&M 端（设备集群 3D 调度）：")

add_heading("1.3 机器人（设备接入）功能维度", level=2)
add_para("本维度定义“让硬件能连进来”的接入标准与适配层。我方不实现硬件本身，但需提供统一接口与协议规范。",
         size=10, color=GREY, space_after=4)
add_bullet("定义统一的设备状态/指令数据契约（JSON 结构），使不同厂商设备即插即用，是接入层的基础。",
           bold_lead="统一接入协议与数据契约：")
add_bullet("以适配器模式对接各厂商设备（ROS2 节点或 MQTT 客户端），把异构硬件统一抽象为“设备代理（DeviceAgent）”。",
           bold_lead="设备驱动适配：")
add_bullet("设备按固定节拍上报心跳与遥测（位置、电量、任务态、健康），后端据此判断在线与可调度性。",
           bold_lead="心跳与遥测上报：")
add_bullet("设备通过 SLAM 或二维码定位得到坐标，上报后端，是地图与路径规划的前提。",
           bold_lead="定位能力：")
add_bullet("设备收到控制指令后回报执行确认与结果，形成“派单→执行→反馈”闭环。",
           bold_lead="指令执行与回执：")
add_bullet("设备监听全局急停等安全主题，软硬结合立即停机，作为演示安全兜底。",
           bold_lead="急停与安全链路：")
add_bullet("深度相机/激光雷达产生真实点云，摄像头提供真实视频流与 AI 识别数据源，驱动 3D 大屏与 PM 端 CCTV。",
           bold_lead="点云与视觉真实源：")

# ============ 1.4 机器人功能实现细则（补充《机器人功能模块规格说明》1.4–1.6） ============
add_heading("1.4 机器人功能实现细则（任务执行 / 状态机 / 安全响应）", level=2)
add_para("本节对《机器人功能模块规格说明》1.4–1.6 三个模块（任务执行引擎、设备状态机、安全响应）补充分实现层面的"
         "四项说明：实现方式、技术类型、关键技术点、落地路径。侧重“怎么做”，为从开发到上线提供可执行指引。"
         "其中 1.4.3 安全响应细分为急停响应、围栏自检、避障、碰撞联锁四个协同子功能。",
         size=10, color=GREY, space_after=6)

add_heading("1.4.1 任务执行引擎", level=3)
add_table(
    ["维度", "内容"],
    [
        ["实现方式",
         "设备收到 assign_task 后，依据下发的 map_point 与 params 做本地路径/动作规划；调度侧 SchedulerEngine "
         "维护任务状态机（pending→running→completed/failed），消费设备回传的 task_event（phase/progress）写入 "
         "Task.progress 与 completed_qty；以 completed_qty≥deliverable_qty 判定完工，达标后解锁 DAG 下游依赖。"],
        ["技术类型",
         "状态机驱动的事件处理（事件溯源）+ 请求-响应（指派/回执）；属“事件驱动 / 流处理”类，进度靠持续"
         "事件流推进而非批处理。"],
        ["关键技术点",
         "① DAG 依赖解锁：用 dependencies（JSONB 任务 ID 列表）表达前置，调度 tick 检查前置 completed 才解锁；"
         "② 进度量化：deliverable_qty/completed_qty 支撑交付量可视；③ 回执幂等：task_event 带 task_id+phase，"
         "重复上报去重；④ 断点续传：设备重连从 started_at 续算，避免重复执行；"
         "⑤ 通信：下行 ROS2 action 或 MQTT QoS1，上行 telemetry 回传进度。"],
        ["落地路径",
         "1) 确认 Task 模型 deliverable_qty/completed_qty/dependencies 解析就绪；"
         "2) SimulatedAdapter.assign_task 实现“派单→模拟规划→周期 emit task_event（开始/进度/完成）”；"
         "3) SchedulerEngine 增加 task_event 消费分支，completed 时解除下游依赖；"
         "4) Ros2Adapter.assign_task 调设备 ROS2 action（如 FollowWaypoints）并订阅 feedback 回填进度；"
         "5) 单台 AGV 跑通“派单→执行→进度→完工→下游解锁”全链后再扩多机；"
         "6) 20 台并发派单 + 每秒进度上报压测，验证 TimescaleDB 连续聚合吞吐。"],
    ],
    widths=[2.8, 13.2],
)

add_heading("1.4.2 设备状态机", level=3)
add_table(
    ["维度", "内容"],
    [
        ["实现方式",
         "以有限状态机统一刻画设备生命周期（idle/moving/working/charging/fault/maintenance/occupied）。"
         "状态存 Device.status；跃迁由事件（指令、task_event、心跳丢失、告警）触发，每次跃迁写 "
         "DeviceEvent(type=status_change) 并发布 Redis 频道→WebSocket 推大屏；跃迁守卫校验合法性"
         "（如 charging 需先达 charging_threshold 才允许→working）。"],
        ["技术类型",
         "有限状态机（FSM）+ 事件溯源；属“状态管理 / 规则引擎”类。"],
        ["关键技术点",
         "① 合法跃迁表（二维矩阵），非法跃迁拒绝并记录；② 跃迁原子性：状态更新+事件写入+推送同事务/异步任务，"
         "避免中间态不一致；③ 体检联动：health（连/定/电/任/安）随状态刷新（fault→safety=fail）；"
         "④ 离线判定：心跳超时（间隔×N）标记 offline+触发告警；"
         "⑤ 前端映射：O&M 5 项体检与状态分布环形图直读 status/health。"],
        ["落地路径",
         "1) Device 模型固化 status 枚举与跃迁守卫（core/state_machine.py）；"
         "2) 适配器每次状态变化调 transition()，统一写事件+推送；"
         "3) SimulatedAdapter 按剧本驱动 idle→moving→working→idle，验证大屏状态分布；"
         "4) 后台任务扫描 last_heartbeat 超时置 offline+告警；"
         "5) Ros2Adapter 订阅设备 /state 话题映射到 FSM；"
         "6) O&M 设备卡片状态标签、PM 资源负载随 FSM 实时刷新联调。"],
    ],
    widths=[2.8, 13.2],
)

add_heading("1.4.3 安全响应（急停 / 围栏自检 / 避障 / 碰撞联锁）", level=3)
add_para("安全响应是系统安全底线，包含四个协同子功能，逐项给出实现方式、技术类型、关键技术点与落地路径：",
         size=10, color=GREY, space_after=4)
add_table(
    ["子功能", "实现方式", "技术类型", "关键技术点", "落地路径"],
    [
        ["急停响应",
         "全局急停按钮触发 send_command(None,'estop')→适配器向所有在线设备广播急停主题/ROS2 服务；"
         "设备立即切断执行器并使能硬件刹车并回执 ack；单机急停按 device_id 下发；急停写 Alert(critical,estop) 推大屏。",
         "发布/订阅广播 + 命令模式；安全关键实时控制（低延迟、QoS 可靠）。",
         "① 双通道：软件急停+硬件急停互为冗余；② 广播寻址 device_id=None 表全量；"
         "③ ACK 超时未确认标记异常告警；④ 恢复门控：estop 后须 reset+人工确认才允许 resume。",
         "1) SimulatedAdapter 实现 estop→置 fault+health.safety=fail+回执；"
         "2) O&M 工具栏急停调 POST /api/devices/command；"
         "3) 告警引擎产生 critical estop alert；"
         "4) Ros2Adapter 调 /emergency_stop 或发布急停主题；"
         "5) estop→reset→resume 回环联调；6) 纳入演示安全剧本。"],

        ["围栏自检",
         "以标段电子围栏（GeoJSON 多边形）为安全边界；设备接收目标/路径前做落点校验，"
         "若落入禁行区/跨标段无授权则拒绝执行并告警；围栏数据来自地图模块区域划分（作业区/禁行区/堆放区/停靠区）。",
         "空间几何计算（点/线段-多边形相交）+ 规则引擎（权限校验）；地理空间 / 约束校验。",
         "① 围栏 GeoJSON+标段 ID；② 射线法点在内 + 线段相交判定 O(顶点数)；"
         "③ 权限矩阵 section_tags vs 区域 section_id，跨区需白名单；"
         "④ 接入已定义但零引用的 device_safety_distance；⑤ 派单前（调度侧）提前拦截而非执行时才发现。",
         "1) 地图模块固化区域模型（禁行区/标段围栏）；"
         "2) 新增 core/geofence.py：is_point_allowed / is_path_allowed；"
         "3) SchedulerEngine 派单前调用，越界标记 failed+告警不派单；"
         "4) SimulatedAdapter 模拟越界验证拦截；"
         "5) 设备端本地围栏校验双保险，越界拒执行并上报；"
         "6) O&M 围栏可视化 + 越界告警。"],

        ["避障",
         "设备基于实时位姿与传感器（激光雷达/深度相机）感知动态障碍（人员、临时堆放），"
         "局部规划层动态避障（减速/绕行）；系统侧通过地图动态障碍标记 + 设备遥测位姿做全局态势感知。",
         "局部路径规划（实时控制）+ 感知-决策闭环；实时控制 / 避障算法。",
         "① 设备端 ROS2 Nav2 局部规划器 + 代价地图动态层；"
         "② 动态障碍数据源：地图标记 + 设备 self 上报 obstacle 列表；"
         "③ 降级：通信丢失切本地避障，不依赖中心；"
         "④ 与碰撞联锁协同：避障失败兜底到联锁（减速/停车）。",
         "1) 地图模块支持动态障碍标记并实时发布；"
         "2) SimulatedAdapter 引入障碍事件驱动局部降速；"
         "3) Ros2Adapter 接入 Nav2 订阅 /scan、/obstacle_layer 回填状态；"
         "4) 安全模块订阅位姿，人员临近安全距离触发降速；"
         "5) 模拟人员穿越路径→设备降速/绕行→大屏告警联调。"],

        ["碰撞联锁",
         "设备间实时交换位姿（中心 Redis 或 ROS2 话题），计算两两距离；距离 < device_safety_distance"
         "（默认 2.0m）触发联锁：相关设备减速/停车，上报联锁事件到监控与安全模块；可分级（预警→降速→停车）。",
         "分布式近邻检测（N-body 距离计算）+ 发布/订阅；实时协同 / 安全联锁。",
         "① 位姿广播：每台周期发布位姿，中心维护设备位姿表；"
         "② 距离矩阵 O(n²)（n≤100 可行）或空间哈希降复杂度；"
         "③ 分级：<0.5×safety→停车，0.5~1×→降速；"
         "④ 缺口：device_safety_distance 已定义但零引用，需新建联锁服务消费位姿流；"
         "⑤ 死锁避免：多机互锁按优先级/先到先得释放，避免集体僵死。",
         "1) 新增 services/interlock.py 订阅位姿流算距离矩阵触发联锁；"
         "2) SimulatedAdapter 双机靠近验证降速/停车 + 事件上报；"
         "3) 接入 device_safety_distance 配置（已存在）参数可调；"
         "4) Ros2Adapter 用多机 /tf 同步或中心集中计算；"
         "5) 监控将联锁事件转告警推大屏；"
         "6) 双机交汇→联锁→错车→恢复纳入演示剧本。"],
    ],
    widths=[1.8, 4.2, 2.6, 4.2, 3.2],
)
doc.add_paragraph()

# ===================== 板块二：技术选型 =====================
doc.add_page_break()
add_heading("板块二 · 技术选型", level=1)
add_para("技术选型在宏观层面确定架构方向与技术策略。下表给出各关键点的选型方向与概念性理由；"
         "其后的说明进一步解释每个选型“是什么、为什么这样选”。", size=10.5, color=GREY, space_after=8)

add_table(
    ["能力域", "选型方向", "核心策略"],
    [
        ["后端服务框架", "FastAPI（异步 Python）", "高并发实时 API + 原生异步，契合调度后台循环"],
        ["实时通信", "WebSocket + Redis 发布订阅", "服务端主动推送，大屏“物理发生什么同步什么”"],
        ["数据持久化", "PostgreSQL + SQLAlchemy；时序用 TimescaleDB", "关系型保主数据，时序库保遥测/事件溯源"],
        ["调度范式", "集中式 + 分布式混合", "全局最优派单 + 局部实时协商避碰"],
        ["派单策略", "市场机制（招标竞价），接口可切强化学习", "可解释、易演示，预留数据驱动优化"],
        ["任务建模", "DAG 有向无环图", "表达 29 项工序的前后依赖与约束"],
        ["设备抽象", "适配器模式（DeviceAgent）", "异构硬件统一抽象，调度核心不改"],
        ["前端框架", "React + TypeScript + Vite", "组件化大屏、类型安全、快热更新"],
        ["状态/样式", "轻量 store（zustand）+ Tailwind CSS", "低样板、原子化样式，适配深色科技风"],
        ["3D 可视化", "Three.js / react-three-fiber", "浏览器内三维沙盘与设备实时标记"],
        ["图表", "轻量自绘 SVG（按需引图表库）", "甘特/趋势/拓扑自绘，零额外依赖"],
        ["机器人接入", "ROS2 为主 + MQTT 桥接", "ROS2 对接智能体节点，MQTT 对接轻量节点"],
        ["定位", "SLAM / 二维码", "沙盘尺度易复现的定位方案"],
        ["报表导出", "Excel（openpyxl）+ PDF（报告引擎）", "满足标准化汇报的多格式导出"],
        ["部署", "Docker Compose", "一键拉起后端+数据库+缓存，便于演示"],
    ],
    widths=[3.2, 5.2, 7.6],
)

doc.add_paragraph()
add_para("选型概念解释：", bold=True, color=NAVY, size=11, space_after=4)
add_bullet("以异步 Python 框架承载 REST 与 WebSocket，天然适配“一个常驻调度循环 + 大量并发设备连接”的模型，"
           "且 Python 在调度算法与 ROS2 生态上同语言、集成成本低。", bold_lead="FastAPI 异步框架：")
add_bullet("实时性是大屏的生命线。调度层状态变化经 Redis 频道发布，再由 WebSocket 透传到前端，"
           "做到“沙盘发生什么、大屏同步什么”，避免前端轮询的延迟与开销。", bold_lead="WebSocket + Redis 发布订阅：")
add_bullet("设备档案、任务、地图等主数据用关系库保证一致与关联；海量遥测/事件用时序库按时间高效写入与聚合，"
           "支撑“事件溯源→多维报表”。", bold_lead="PostgreSQL + 时序库：")
add_bullet("全局派单追求“全流程有序推进”，由集中式求解；局部资源争用与避碰追求“实时性”，由分布式协商。两者解耦协作，"
           "兼顾全局最优与局部响应。", bold_lead="集中式 + 分布式混合调度：")
add_bullet("把任务当作“标的”，让设备按成本（距离、电量、负载）投标，调度器定标。规则透明、演示直观，"
           "且通过策略接口预留切换到强化学习（从历史演示学最优派单）。", bold_lead="市场机制派单：")
add_bullet("用有向无环图表达工序依赖，只有前置完成的任务才“解锁”可派，从数据结构上保证施工顺序正确。",
           bold_lead="DAG 任务建模：")
add_bullet("为每类厂商设备写一层“适配器”，向上暴露统一接口。新增设备类型只需加适配器，调度核心与前端零改动。",
           bold_lead="适配器模式（DeviceAgent）：")
add_bullet("React 组件化适合多面板大屏；TypeScript 在 8 大模块、双端共享类型下显著降低联调错误；"
           "Vite 提供极快的热更新，提升大屏迭代效率。", bold_lead="React + TS + Vite：")
add_bullet("zustand 以极简 API 管理跨面板共享状态；Tailwind 原子化类名快速搭出深色科技风界面，避免样式文件膨胀。",
           bold_lead="zustand + Tailwind：")
add_bullet("浏览器内 WebGL 三维渲染，直接把点云与设备坐标映射为可旋转、可缩放的沙盘，无需插件。",
           bold_lead="Three.js / react-three-fiber：")
add_bullet("甘特图、趋势线、拓扑等用 SVG 手绘即可满足，省去重型图表库；若后续需复杂交互再引入专业图表库。",
           bold_lead="轻量自绘 SVG：")
add_bullet("ROS2 是机器人智能体的事实标准，便于对接各设备厂商节点；MQTT 轻量、适合沙盘低功耗/微控制器节点，"
           "二者通过桥接共存，覆盖不同档位设备。", bold_lead="ROS2 + MQTT：")
add_bullet("Excel 满足可编辑明细，PDF 满足正式汇报，两者并行覆盖“分析”与“汇报”两类场景。",
           bold_lead="Excel + PDF 报表：")

# ===================== 板块三：技术路线 =====================
doc.add_page_break()
add_heading("板块三 · 技术路线", level=1)
add_para("技术路线回答“按什么顺序做”。以下在方案既定 M1–M5 里程碑基础上，叠加当前缺口的优先级（P0/P1/P2），"
         "形成可执行的分阶段路径。", size=10.5, color=GREY, space_after=8)

add_table(
    ["阶段", "目标", "关键交付物", "依赖与优先级"],
    [
        ["M1 接入与地图", "打通设备接入与空间基础", "设备代理框架、注册表、网格+点位地图", "基础；建议补齐真实 Ros2Adapter（P0）"],
        ["M2 调度内核", "可离线运行的调度引擎", "任务 DAG、市场机制派单、批量派单", "依赖 M1；当前已完成度最高"],
        ["M3 协同与安全", "规模化协同与安全闭环", "群组指令、编队、碰撞联锁、急停、电子围栏", "依赖 M2；碰撞联锁/围栏为 P0"],
        ["M4 可视化与报表", "大屏与数据出口", "数字孪生大屏、告警引擎、多维报表导出", "依赖 M1–M3；告警引擎为 P0，PDF 为 P1"],
        ["M5 联调演示", "硬件联调与定型", "可演示系统、演示剧本、培训与文档、Alembic/测试", "依赖全部；工程化为 P2"],
    ],
    widths=[3.0, 4.2, 5.0, 4.0],
)

doc.add_paragraph()
add_para("缺口优先级叠加（基于当前实现 vs 方案）：", bold=True, color=NAVY, size=11, space_after=4)
add_bullet("演示/验收阻断项，必须最先补齐：① 真实机器人接入层（ROS2/MQTT 适配器）；"
           "② 告警生成引擎（让大屏“看得见异常”）；③ 碰撞联锁 + 电子围栏越界拦截（安全底线）。",
           bold_lead="P0：")
add_bullet("展示完整度：O&M 三维透明方块模型与三视图真实切换；续航预测 + 维保提醒；"
           "摄像头真实视频流 + AI 识别数据源；PDF 报表导出。", bold_lead="P1：")
add_bullet("工程化与长期可维护：Alembic 数据库迁移、自动化测试；分布式调度/协商、任务模板、"
           "动态障碍、历史告警/轨迹回放、待办多分类标签等体验增强。", bold_lead="P2：")

doc.add_paragraph()
hrule()
add_para("一句话总结：后端“大脑逻辑”已成形、前端双端骨架已立，但从“能跑的演示原型”到“方案承诺的科研样机”，"
         "最关键的真实设备接入、安全执行闭环、告警引擎三块仍是空白，应沿 M1→M5 路线优先投入 P0 项。",
         size=10.5, color=NAVY, bold=True, space_after=4)

doc.save(OUT)
print("SAVED:", OUT)
