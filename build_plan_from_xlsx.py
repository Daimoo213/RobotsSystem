# -*- coding: utf-8 -*-
"""将 技术方案总览（修改）.xlsx 的功能点写入 方案一_集群调度系统开发方案.docx 的第五章。"""
import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EXCEL = r"D:\MengWork\Web\RobotsSystem\技术方案总览（修改）.xlsx"
SRC = r"D:\MengWork\Web\RobotsSystem\方案一_集群调度系统开发方案.docx"
OUT = r"D:\MengWork\Web\RobotsSystem\方案一_集群调度系统开发方案_功能点版.docx"

EA = "微软雅黑"
STATUS_COLOR = {
    "已实现": RGBColor(0x2E, 0x7D, 0x32),
    "部分实现": RGBColor(0xC8, 0x6A, 0x00),
    "未实现": RGBColor(0xC0, 0x39, 0x2B),
}
STATUS_LABEL = {"已实现": "已落地", "部分实现": "部分落地", "未实现": "待建设"}

# ---------- 读取 Excel ----------
wb = openpyxl.load_workbook(EXCEL)
ws = wb["功能·技术选型总表"]
fps = []
cur = None
r = 3
while r <= ws.max_row:
    dim = ws.cell(r, 1).value
    num = ws.cell(r, 2).value
    fp = ws.cell(r, 3).value
    d = ws.cell(r, 4).value
    tt = ws.cell(r, 5).value
    sel = ws.cell(r, 6).value
    kt = ws.cell(r, 7).value
    rt = ws.cell(r, 8).value
    st = ws.cell(r, 9).value
    if fp:
        cur = {"dim": dim, "num": num, "name": fp, "status": st, "subs": []}
        fps.append(cur)
    if d:
        cur["subs"].append((d, tt, sel, kt, rt))
    r += 1

ws2 = wb["技术点介绍"]
techpoints = []
for rr in range(3, ws2.max_row + 1):
    row = [ws2.cell(rr, c).value for c in range(1, 5)]
    if any(v for v in row):
        techpoints.append(row)

# 按维度分组
def by_dim(prefix):
    return [f for f in fps if f["num"].startswith(prefix)]

BACKEND = by_dim("B")
FRONTEND = by_dim("F")
DATABASE = by_dim("D")
print(f"后端 {len(BACKEND)} 前端 {len(FRONTEND)} 数据库 {len(DATABASE)} 技术点 {len(techpoints)}")


# ---------- 工具函数 ----------
def set_cjk(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), EA)
    rFonts.set(qn('w:ascii'), EA)
    rFonts.set(qn('w:hAnsi'), EA)


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def new_para(doc, text, style=None, bold=False, size=11, color=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    set_cjk(run)
    el = p._p
    doc.element.body.remove(el)
    return el


def new_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = 1  # center
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        set_cjk(run)
        set_cell_bg(hdr[i], "D9E2F3")
        hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val) if val is not None else "")
            run.font.size = Pt(9)
            set_cjk(run)
            cells[i].width = Cm(widths[i])
    el = t._tbl
    doc.element.body.remove(el)
    return el


# ---------- 打开原文档，定位第五章 ----------
doc = Document(SRC)
body = doc.element.body

def find_heading(text, level="Heading 1"):
    for p in doc.paragraphs:
        if p.style.name == level and p.text.strip() == text:
            return p._p
    return None

h5 = find_heading("五、集群调度系统功能模块详细实现")
h6 = find_heading("六、模块间逻辑关系（控制流 / 数据流总览）")
assert h5 is not None, "未找到第五章标题"
assert h6 is not None, "未找到第六章标题"

# 删除第五章原有内容（h5 与 h6 之间）
to_remove = []
between = False
for el in list(body):
    if el is h5:
        between = True
        continue
    if el is h6:
        break
    if between:
        to_remove.append(el)
for el in to_remove:
    body.remove(el)

# 重命名第五章标题
for run in h5.findall(qn('w:p') + '/w:r') if False else []:
    pass
# 直接改 run 文本
_runs = h5.findall('.//' + qn('w:r'))
for run in _runs:
    t = run.find(qn('w:t'))
    if t is not None:
        t.text = "五、系统功能模块详细实现"
        for rn in run.findall(qn('w:rPr') + '/' + qn('w:rFonts')):
            rn.set(qn('w:eastAsia'), EA)

# ---------- 构建第五章新内容 ----------
new_els = []

new_els.append(new_para(doc,
    "本章以《技术方案总览》之“功能·技术选型总表”的功能点为基准，从后端、前端、数据库三个维度对系统功能模块进行细化定义。"
    "每个功能点下设若干子能力，分别标注技术类型、技术选型（或真实做法）、关键技术点与技术路线，并给出实现状态，"
    "作为后续开发、验收与排期的统一依据。", size=10.5))

# 维度章节
def dim_block(title, fps_list, chap, intro):
    els = []
    els.append(new_para(doc, title, style="Heading 2"))
    els.append(new_para(doc, intro, size=10.5))
    for idx, fp in enumerate(fps_list, 1):
        sec = f"{chap}.{idx}"
        head = f"{sec} 〔{fp['num']}〕{fp['name']}"
        els.append(new_para(doc, head, style="Heading 3"))
        st = fp["status"] or "未实现"
        label = STATUS_LABEL.get(st, st)
        els.append(new_para(doc, f"实现状态：{label}（{st}）",
                            bold=True, size=10, color=STATUS_COLOR.get(st, RGBColor(0,0,0))))
        rows = [(d, tt, sel, kt, rt) for (d, tt, sel, kt, rt) in fp["subs"]]
        els.append(new_table(doc,
            ["功能定义", "技术类型", "技术选型（真实做法）", "关键技术点", "技术路线（落地步骤）"],
            rows, [3.0, 2.0, 3.6, 3.4, 4.0]))
    return els

new_els += dim_block("5.1 后端功能模块", BACKEND, "5.1",
    "后端为调度系统“大脑”，承载设备接入、地图感知、智能调度、协同控制、能耗维保、监控告警、数据报表与安全管控八大模块。")
new_els += dim_block("5.2 前端功能模块", FRONTEND, "5.2",
    "前端包含 PM 端（施工进度指挥）与 O&M 端（设备集群 3D 调度）两套大屏，二者共用同一后端，以下按功能点分别定义。")
new_els += dim_block("5.3 数据库功能模块", DATABASE, "5.3",
    "数据库层提供业务关系存储、实时缓存与发布订阅、时序数据落库、版本化迁移与备份高可用能力。")

# 插入到 h6 之前
h6_idx = list(body).index(h6)
for i, el in enumerate(new_els):
    body.insert(h6_idx + i, el)

# ---------- 附录：技术点说明 ----------
sectPr = body.find(qn('w:sectPr'))
append_els = []
append_els.append(new_para(doc, "附录一、关键技术点说明", style="Heading 1"))
append_els.append(new_para(doc,
    "本附录对应《技术方案总览》之“技术点介绍”表，对方案中涉及的关键技术名词给出概念解释与本项目用途，"
    "便于团队快速理解整体技术方案。", size=10.5))
append_els.append(new_table(doc,
    ["技术点", "类别", "概念解释", "本项目用途"],
    techpoints, [2.6, 2.2, 6.0, 5.2]))
for i, el in enumerate(append_els):
    sectPr.addprevious(el)

# ---------- 保存 ----------
doc.save(OUT)
print("SAVED:", OUT)
