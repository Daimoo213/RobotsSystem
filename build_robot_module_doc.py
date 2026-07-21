# -*- coding: utf-8 -*-
"""生成《机器人功能模块规格说明》DOCX，补充技术方案总览的机器人（设备接入）层。"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

CJK = "Microsoft YaHei"

doc = Document()

# 默认正文字体（含中文）
normal = doc.styles["Normal"]
normal.font.name = CJK
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)


def set_cjk(run):
    run.font.name = CJK
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), CJK)


def h1(text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    set_cjk(run)
    return p


def h2(text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    set_cjk(run)
    return p


def h3(text):
    p = doc.add_heading(level=3)
    run = p.add_run(text)
    set_cjk(run)
    return p


def para(text, bold=False, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_cjk(run)
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    run = p.add_run(text)
    set_cjk(run)
    return p


def num(text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_cjk(run)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(9.5)
        set_cjk(run)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
            set_cjk(run)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


# ============================================================
# 封面
# ============================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
trun = title.add_run("机器人功能模块规格说明")
trun.bold = True
trun.font.size = Pt(22)
set_cjk(trun)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
srun = sub.add_run("—— 设备接入与集群管理层（技术方案总览 · 机器人维度补充）")
srun.font.size = Pt(12)
srun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
set_cjk(srun)

doc.add_paragraph()

# ============================================================
# 0. 文档说明与术语映射
# ============================================================
h1("0. 文档说明与术语映射")

para("本文档是《技术方案总览（功能维度 / 技术选型 / 技术路线）》的机器人（设备）维度专项补充，"
     "面向设备接入与集群管理这一核心模块，系统梳理机器人侧的核心功能、配置参数与接口，以及与现有"
     "调度系统的集成方式。遵循原方案交付边界：本系统交付调度系统前后端与标准化设备接入 API，"
     "不涉及机器人硬件本体的控制算法实现。")

para("术语映射说明：", bold=True)
para("为避免歧义，将本次需求中偏“对话式”的表述映射到本项目的施工机器人语境：")
table(
    ["需求表述", "在本项目中的对应含义"],
    [
        ["自动回复", "设备对调度指令的自动应答（ACK/NACK）与主动遥测自报（心跳、位姿、电量），"
                    "即“无需轮询、设备主动上报状态并确认每条指令”。"],
        ["消息处理", "接入层（DeviceAgent 适配器）对上行/下行消息的统一处理管线：接收→解析→校验→路由→响应。"],
        ["任务执行", "设备端接收派单后，从接单、本地规划、执行到进度回传与完工确认的闭环能力。"],
    ],
    widths=[1.6, 4.9],
)

# ============================================================
# 1. 核心功能定义与实现逻辑
# ============================================================
h1("1. 核心功能定义与实现逻辑")

h2("1.1 设备注册与一机一档")
para("功能定义：每台机器人首次接入时，自动完成身份注册并建立“一机一档”档案，作为后续所有调度、"
     "运维、安全管控的数据基座。")
bullet("自动注册：设备上线即广播注册信息，调度中心自动发现并写入设备注册表（内存 + 持久化）。")
bullet("能力声明：设备上报自身型号、类型（agv/excavator/crane/masonry/inspect）、可执行工序与传感能力，"
       "写入 capabilities 字段。")
bullet("身份与权限：依据 section_tags（标段）、permissions（通行/施工权限）建立三维分组（标段/车型/工序），"
       "作为调度派单与安全分区的约束依据。")
bullet("档案持久化：注册时间、最近心跳时间（last_heartbeat）等纳入设备主档，支撑在线率统计与体检指标。")

h2("1.2 指令自动应答与遥测自报（对应“自动回复”）")
para("功能定义：设备侧以“事件驱动 + 周期自报”方式，主动向系统反馈状态，而非被动等待轮询。这是"
     "数字孪生大屏“物理沙盘发生什么、大屏同步什么”的数据来源。")
bullet("心跳自报：按固定间隔上报 last_heartbeat，超时未报即判定离线，触发在线率下降与告警。")
bullet("指令应答（ACK/NACK）：每收到一条控制指令，设备须在限定时间内回执确认（成功/拒绝及原因），"
       "调度侧据此判断指令是否生效。")
bullet("状态变化上报：设备状态机发生跃迁（如 idle→moving、working→fault）时即时上报，"
       "驱动大屏状态分布与体检指标（连/定/电/任/安）刷新。")
bullet("遥测自报：周期性上报位姿（position_x/y/z）、电量（battery）、当前任务进度等，"
       "供地图感知、能耗预测、监控中心消费。")

h2("1.3 消息处理管线（对应“消息处理”）")
para("功能定义：接入层对全部上行（设备→系统）与下行（系统→设备）消息执行统一的五阶段处理，"
     "保证数据一致、可追溯、可路由。")
num("接收：通过 ROS2 主题或 MQTT 主题订阅设备消息，按设备 ID 分流。")
num("解析：将二进制/JSON 报文反序列化为统一内部消息结构。")
num("校验：校验设备身份、消息完整性、字段合法性及安全约束（如目标点是否在禁行区）。")
num("路由：依据消息类别分发至对应处理单元——注册消息入注册表、遥测消息入状态/地图、"
    "告警消息入告警引擎、任务回执入调度器。")
num("响应：对需要应答的消息生成 ACK/NACK 或状态回执，按原通道回送设备。")
para("下行命令类型（已实现接口契约）：pause（暂停）、resume（恢复）、estop（急停）、"
     "reset（重置）、release（释放）、assign_task（派单）。", bold=True)

h2("1.4 任务执行引擎（对应“任务执行”）")
para("功能定义：设备接收调度派单后，在本地完成从接单到完工的闭环，并向系统持续回传执行状态。")
num("接单：收到 assign_task，校验自身状态与能力是否满足，回执 ACK/NACK。")
num("本地规划：依据目标点位（map_point）与当前位姿进行路径/动作规划（硬件侧或轻量规划器完成）。")
num("执行：按规划动作推进，期间持续上报进度（progress）与位姿。")
num("进度回传：以 task_event 形式上报“开始/进行中/完成/异常”，进度同步写入任务主档。")
num("完工确认：达成交付量（completed_qty）后上报 completed，调度器据此推进 DAG 下游工序。")

h2("1.5 设备状态机")
para("功能定义：以有限状态机统一刻画设备生命周期，作为派单可行性判断与体检指标的基础。")
table(
    ["状态", "含义", "典型跃迁触发"],
    [
        ["idle 待机", "空闲可调度", "收到派单→working；人工指令→moving/charging"],
        ["moving 移动", "前往任务点途中", "抵达→working；急停→fault/stopped"],
        ["working 作业", "执行任务中", "完工→idle；异常→fault"],
        ["charging 充电", "错峰充能中", "达充电阈值→idle"],
        ["fault 故障", "异常/离线", "人工复位→idle"],
        ["maintenance 维护", "保养/检修工单", "完成→idle"],
        ["occupied 占用", "被锁定/占用", "释放→idle"],
    ],
    widths=[1.3, 2.4, 2.8],
)

h2("1.6 安全响应")
bullet("急停响应：收到 estop 指令后设备须立即停机并回执，全局急停对所有在线设备广播。")
bullet("围栏自检：移动前校验目标点与路径是否越界（禁行区/跨标段无授权），越界则拒绝执行并上报。")
bullet("避障与碰撞联锁：实时交换邻机位姿，进入 device_safety_distance（默认 2.0 米）触发减速/停车，"
       "并将联锁事件上报监控与安全模块。")

# ============================================================
# 2. 配置参数与接口说明
# ============================================================
h1("2. 配置参数与接口说明")

h2("2.1 设备配置参数表（DeviceAgent）")
para("以下参数对应后端 Device 模型与 config.py 中的真实字段，作为设备档案与运行期配置的统一来源。")
table(
    ["参数", "类型", "默认/示例", "说明"],
    [
        ["code / name", "string", "T-03 / 03号运输车", "设备唯一编码与显示名"],
        ["type", "enum", "agv/excavator/crane/masonry/inspect", "设备类型，决定可派工序"],
        ["model", "string", "厂商型号", "硬件型号，用于维保匹配"],
        ["capabilities", "dict(JSONB)", "{}", "能力声明（工序/传感）"],
        ["section_tags", "dict(JSONB)", "{}", "所属标段分组标签"],
        ["permissions", "dict(JSONB)", "{}", "通行/施工权限"],
        ["status", "enum", "idle", "当前状态机取值"],
        ["battery", "float", "100.0", "电量百分比"],
        ["position_x/y/z", "float", "0.0", "实时位姿（米）"],
        ["section_id", "string", "标段ID", "当前所在标段"],
        ["last_heartbeat", "datetime", "—", "最近心跳时间，离线判定依据"],
        ["health", "dict(JSONB)", "{}", "体检指标：connection/location/battery/task/safety"],
        ["heartbeat_interval", "float", "依型号", "心跳自报周期（秒）"],
        ["low_battery_threshold", "float", "20.0", "低电量阈值，触发降权/充电"],
        ["charging_threshold", "float", "90.0", "充电停止阈值"],
        ["device_safety_distance", "float", "2.0（米）", "碰撞联锁触发距离"],
        ["geofence", "geo", "标段电子围栏", "分区管控边界"],
        ["comm_protocol", "enum", "ros2 / mqtt", "通信协议选型"],
    ],
    widths=[1.7, 1.2, 1.6, 2.0],
)

h2("2.2 标准接入接口")
para("架构方向：采用“适配器模式 + 统一协议契约”。系统定义与硬件无关的 DeviceAdapter 抽象，"
     "真实设备通过 Ros2Adapter（ROS2 节点）或 MQTT 轻量节点接入，仿真设备由 SimulatedAdapter 承载，"
     "三者对外暴露一致接口，调度层无感知。")
table(
    ["接口类别", "方向", "约定（概念级）", "说明"],
    [
        ["注册", "上行", "REST/主题 注册消息", "设备上线写入注册表"],
        ["心跳", "上行", "周期主题 heartbeat", "在线判定与体检"],
        ["遥测", "上行", "主题 telemetry", "位姿/电量/进度"],
        ["指令", "下行", "主题/服务 command", "pause/resume/estop/reset/release"],
        ["派单", "下行", "主题/服务 assign_task", "携带任务与 map_point"],
        ["任务回执", "上行", "主题 task_event", "开始/进度/完成/异常"],
        ["告警", "上行", "主题 alert", "越界/低电/故障等"],
    ],
    widths=[1.2, 0.8, 2.2, 2.3],
)

h2("2.3 消息数据模型（概念字段）")
table(
    ["消息", "关键字段（概念）", "用途"],
    [
        ["RegistrationMsg", "device_id, type, model, capabilities, section_tags, permissions", "建立一机一档"],
        ["HeartbeatMsg", "device_id, timestamp, status", "在线判定"],
        ["TelemetryMsg", "device_id, pos_x/y/z, battery, progress, health", "态势与能耗输入"],
        ["CommandMsg", "device_id|null, cmd(pause/estop/...), ts", "下发控制"],
        ["TaskAssignMsg", "device_id, task_id, map_point, params", "派单"],
        ["TaskStatusMsg", "device_id, task_id, phase, progress, qty", "执行回执"],
        ["AlarmMsg", "device_id, level, category, message, payload", "告警产生"],
    ],
    widths=[1.5, 3.0, 2.0],
)

h2("2.4 模块间数据交互")
para("设备层是全部机器人数据的入口，向各模块单向/双向供给数据：")
bullet("→ 地图感知：实时坐标（position_x/y/z）驱动沙盘镜像与动态障碍更新。")
bullet("→ 智能调度：可调度资源清单（idle 设备 + capabilities）与约束（电量/标段权限）。")
bullet("→ 能耗维保：电量与里程，支撑续航预测、错峰充能与维保提醒。")
bullet("→ 安全管控：在线状态与位姿，支撑电子围栏与碰撞联锁约束。")
bullet("← 调度/安全/人工：下行指令经消息管线回送设备，形成“感知—决策—执行—反馈”闭环。")
bullet("→ 监控告警 / 报表：全部状态、事件、告警汇聚，驱动大屏与统计沉淀。")

# ============================================================
# 3. 与现有系统的集成方式
# ============================================================
h1("3. 与现有系统的集成方式")

h2("3.1 触发条件")
table(
    ["触发条件", "来源", "系统动作"],
    [
        ["设备上线", "设备广播注册", "自动发现→写入注册表→纳入可调度"],
        ["调度派单", "SchedulerEngine 市场机制定标", "assign_task 下行→设备接单"],
        ["告警阈值", "电量/越界/超时/故障", "告警引擎产生→大屏+处置闭环"],
        ["人工干预", "PM/O&M 界面", "pause/resume/estop/reset/release 下行"],
        ["急停", "全局急停按钮", "对所有在线设备广播 estop"],
    ],
    widths=[1.4, 2.0, 3.1],
)

h2("3.2 调用流程（主流程）")
num("设备上线：广播 RegistrationMsg → 注册表建立档案 → 状态置 idle。")
num("心跳/遥测循环：设备周期自报 → 刷新状态、位姿、体检指标 → 推送大屏。")
num("派单：调度器依据 DAG 与约束选中设备 → assign_task 下行 → 设备 ACK 并置 working。")
num("执行与回传：设备本地规划执行 → task_event 持续回传进度 → 进度写入 Task 主档。")
num("完工：completed_qty 达标 → 上报 completed → 调度器解锁 DAG 下游工序。")
num("下线/异常：心跳丢失或 fault → 标记离线/故障 → 触发告警与重规划。")

h2("3.3 异常处理机制")
table(
    ["异常", "判定", "处理"],
    [
        ["离线（心跳丢失）", "超过心跳超时窗", "标记离线→在线率下降→告警→调度器剔除可调度清单"],
        ["指令超时", "下发后未在 ACK 窗内回执", "重试（有限次）→仍失败标记设备异常→转派其他设备"],
        ["任务失败", "task_event 上报异常", "任务置 failed→调度器重规划/改派→记录干预留痕"],
        ["安全越界", "目标/路径落入禁行区或跨标段无授权", "拒绝执行→上报告警→安全模块拦截派单"],
        ["碰撞风险", "邻机距离 < safety_distance", "减速/停车→上报联锁事件→安全模块复核"],
        ["通信错误", "报文损坏/通道中断", "指数退避重传→断连后进入离线流程→恢复后重新注册"],
    ],
    widths=[1.3, 2.1, 3.1],
)

h2("3.4 与各模块的集成点")
bullet("调度模块：消费 capabilities/status 做派单，输出 assign_task；受能耗/安全约束。")
bullet("安全模块：输入位姿与在线状态，输出禁行/急停/围栏约束，接收越界与联锁事件。")
bullet("监控告警：订阅全部设备事件，告警引擎产生分级告警并关联处置闭环。")
bullet("能耗维保：输入电量/里程，输出充能调度与维保工单。")
bullet("地图感知：输入坐标，输出地图/点位/障碍供调度校验。")
bullet("报表统计：全量事件时序沉淀，反哺调度策略优化。")

# ============================================================
# 4. 实现现状与落地建议
# ============================================================
h1("4. 实现现状与落地建议")

h2("4.1 已实现（后端）")
bullet("Device 数据模型完整：含 code/type/capabilities/permissions/status/battery/位姿/health 等字段。")
bullet("DeviceAdapter 抽象与 SimulatedAdapter：接口契约（start/stop/assign_task/send_command/…）已定义并仿真可用。")
bullet("配置项落地：low_battery_threshold、charging_threshold、device_safety_distance、mqtt_* 等已在 config.py 定义。")
bullet("设备事件模型 DeviceEvent（telemetry/status_change/task_event/alert/environment）已建表。")

h2("4.2 待补齐（方案规划但代码缺失）")
bullet("真实设备接入：Ros2Adapter 未实现，MQTT 未接线（mqtt_enabled=False），仅仿真。")
bullet("消息处理落地：适配器对真实报文的解析/校验/路由链路未接通。")
bullet("碰撞联锁 + 电子围栏拦截：device_safety_distance 已定义但零引用，移动逻辑不校验禁行区。")
bullet("告警产生引擎：仅有查询接口，无产生逻辑，大屏告警恒为空。")
bullet("续航预测、维保提醒、编队同步：尚未实现。")

h2("4.3 落地优先级")
table(
    ["优先级", "事项", "理由"],
    [
        ["P0（演示阻断）", "真实设备接入（Ros2Adapter+MQTT）、告警引擎、碰撞联锁+围栏拦截",
         "方案核心交付物，缺则“能看不能控/不能生”"],
        ["P1", "编队同步、维保提醒、续航预测、任务模板", "提升演示完整度与可持续性"],
        ["P2", "分布式协商调度、动态障碍物、WebSocket 历史回放", "能力深化与体验增强"],
    ],
    widths=[1.3, 3.0, 2.2],
)

doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
frun = foot.add_run("本文档与《技术方案总览（功能维度/技术选型/技术路线）》配套使用，"
                    "为机器人（设备接入）层的专项细化。")
frun.font.size = Pt(9)
frun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
set_cjk(frun)

out = "D:/MengWork/Web/RobotsSystem/机器人功能模块规格说明.docx"
doc.save(out)
print("saved:", out)
