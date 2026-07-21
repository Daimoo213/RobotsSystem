# -*- coding: utf-8 -*-
from docx import Document
path = r"D:/MengWork/Web/RobotsSystem/采购需求_基础设施参数（30机器人）.docx"
doc = Document(path)
print("=== 段落 ===")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)
print("\n=== 技术参数表格 ===")
t = doc.tables[0]
for ri, row in enumerate(t.rows):
    print(f"--- ROW {ri} ---")
    for c in row.cells:
        if c.text.strip():
            print(c.text)
