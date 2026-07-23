from docx import Document
doc=Document(r"D:/MengWork/Web/RobotsSystem/采购需求_基础设施参数（30机器人）.docx")
print("===== PARAGRAPHS =====")
for i,p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("===== TABLES:", len(doc.tables), "=====")
# check merged cells
t=doc.tables[0]
print("rows:", len(t.rows), "cols:", len(t.columns))
for ri,row in enumerate(t.rows):
    for ci,cell in enumerate(row.cells):
        # check if merged
        pass
