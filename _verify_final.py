# -*- coding: utf-8 -*-
from docx import Document
import re
doc = Document(r"D:/MengWork/Web/RobotsSystem/采购需求_基础设施参数（30机器人）.docx")
spec = "\n".join(c.text for row in doc.tables[0].rows for c in row.cells)

# 每个 ▲ 项都应含（投标时提供
ast_items = re.findall(r'▲\d+\.\d+ [^\n]*', spec)
missing_note = [it[:12] for it in ast_items if "（投标时提供" not in it]
print("▲ 总数:", len(ast_items), "| 缺备注的▲:", missing_note if missing_note else "NONE")
print("（投标时提供 出现次数:", spec.count("（投标时提供"))

checks = {
 "1.7降1080p": "支持双屏 1920×1080 输出（投标时提供设备检测报告）。" in spec,
 "无4K": "3840×2160" not in spec,
 "摄像头=4台": "摄像机数量：4 台。" in spec,
 "无功能残留": not any(w in spec for w in ["渲染","调度引擎","数据库","遥测","点云","沙盘","VLAN","QoS","漫游","时延","录像","回放","AI 识别","跨镜头","无盲区","演示","场景","RTSP 拉流","对接大屏","控制网","急停","WebGL","CUDA","三维","微缩","轨迹","融合"]),
 "商务无演示": "演示" not in "\n".join(p.text for p in doc.paragraphs),
 "商务无30台机器人": "30 台机器人" not in "\n".join(p.text for p in doc.paragraphs),
 "商务无ROS2": "ROS2" not in "\n".join(p.text for p in doc.paragraphs),
}
for k,v in checks.items(): print(("OK  " if v else "FAIL ")+k)
print("\nRESULT:", "ALL PASS" if (not missing_note and all(checks.values())) else "CHECK")
